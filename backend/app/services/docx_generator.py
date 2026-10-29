import os
import re
from datetime import datetime
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from app.core.config import settings

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)

def generate_smart_notes_docx(
    course_title: str,
    section_title: str,
    objectives: list,
    content: str,
    summary: str = "",
    output_filename: str = ""
) -> str:
    """
    Generates a professionally formatted Word (.docx) document for Smart Notes.
    Returns the absolute path to the generated document.
    """
    if not output_filename:
        safe_name = re.sub(r'[^a-zA-Z0-9_-]', '_', f"{course_title}_{section_title}")[:40]
        output_filename = f"notes_{safe_name}.docx"
        
    output_path = os.path.join(settings.NOTES_DIR, output_filename)
    
    doc = Document()
    
    # Page Margins
    sections = doc.sections
    for s in sections:
        s.top_margin = Inches(1.0)
        s.bottom_margin = Inches(1.0)
        s.left_margin = Inches(1.0)
        s.right_margin = Inches(1.0)

    # Document Header / Brand
    p_brand = doc.add_paragraph()
    r_brand = p_brand.add_run("EDUAI LEARNING PLATFORM — SMART STUDY NOTES")
    r_brand.font.size = Pt(9.5)
    r_brand.font.bold = True
    r_brand.font.color.rgb = RGBColor(79, 70, 229) # Indigo
    
    # Title
    p_title = doc.add_paragraph()
    r_title = p_title.add_run(course_title)
    r_title.font.size = Pt(22)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(23, 32, 51)
    
    # Section Subtitle
    p_sec = doc.add_paragraph()
    r_sec = p_sec.add_run(f"Section Topic: {section_title}")
    r_sec.font.size = Pt(14)
    r_sec.font.bold = True
    r_sec.font.color.rgb = RGBColor(99, 102, 241)
    
    p_meta = doc.add_paragraph()
    r_meta = p_meta.add_run(f"Academic Date: {datetime.now().strftime('%B %d, %Y')}  |  Curriculum Standard: University Accredited")
    r_meta.font.size = Pt(9)
    r_meta.font.italic = True
    r_meta.font.color.rgb = RGBColor(100, 116, 139)
    
    doc.add_paragraph("—" * 50)
    
    # Summary Callout Box
    if summary:
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = tbl.cell(0, 0)
        cell.width = Inches(6.5)
        set_cell_background(cell, "F1F5F9")
        p_box = cell.paragraphs[0]
        r_box_title = p_box.add_run("Executive Summary: ")
        r_box_title.font.bold = True
        r_box_title.font.size = Pt(10)
        r_box_title.font.color.rgb = RGBColor(30, 41, 59)
        r_box_text = p_box.add_run(summary)
        r_box_text.font.size = Pt(10)
        r_box_text.font.color.rgb = RGBColor(51, 65, 85)
        doc.add_paragraph()

    # Objectives
    if objectives:
        h_obj = doc.add_heading("1. Core Learning Objectives", level=2)
        for r in h_obj.runs:
            r.font.color.rgb = RGBColor(23, 32, 51)
        for obj in objectives:
            doc.add_paragraph(obj, style='List Bullet')
        doc.add_paragraph()

    # Section Content
    h_content = doc.add_heading("2. Detailed Technical & Conceptual Notes", level=2)
    for r in h_content.runs:
        r.font.color.rgb = RGBColor(23, 32, 51)
        
    lines = content.split('\n')
    for line in lines:
        stripped = line.strip()
        if not stripped:
            continue
        if stripped.startswith('### '):
            h = doc.add_heading(stripped[4:], level=3)
            for r in h.runs:
                r.font.color.rgb = RGBColor(79, 70, 229)
        elif stripped.startswith('## '):
            h = doc.add_heading(stripped[3:], level=2)
            for r in h.runs:
                r.font.color.rgb = RGBColor(30, 41, 59)
        elif stripped.startswith('# '):
            pass
        elif stripped.startswith('- ') or stripped.startswith('* '):
            doc.add_paragraph(stripped[2:], style='List Bullet')
        elif 'tip' in stripped.lower() or 'exam' in stripped.lower() or 'note:' in stripped.lower():
            p_tip = doc.add_paragraph()
            r_tip_tag = p_tip.add_run("[ EXAM TIP ] ")
            r_tip_tag.font.bold = True
            r_tip_tag.font.color.rgb = RGBColor(194, 65, 12) # Amber
            r_tip_body = p_tip.add_run(stripped)
            r_tip_body.font.italic = True
        elif stripped.startswith('```'):
            continue
        else:
            doc.add_paragraph(stripped)

    # Checklist
    doc.add_paragraph()
    h_check = doc.add_heading("3. Exam Readiness Checklist", level=2)
    for r in h_check.runs:
        r.font.color.rgb = RGBColor(23, 32, 51)
    
    check_table = doc.add_table(rows=1, cols=2)
    check_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr_cells = check_table.rows[0].cells
    hdr_cells[0].text = "Knowledge Objective"
    hdr_cells[1].text = "Status"
    set_cell_background(hdr_cells[0], "E2E8F0")
    set_cell_background(hdr_cells[1], "E2E8F0")
    
    checklist_items = [
        "Primary Mathematical Model & Loss Objective",
        "Distinction between Supervised, Unsupervised, and RL",
        "Evaluation Metrics (Recall, Precision, ROC-AUC, F1)",
        "Overfitting Mitigation (L1 vs L2 Regularization, Early Stopping)"
    ]
    for item in checklist_items:
        row_cells = check_table.add_row().cells
        row_cells[0].text = item
        row_cells[1].text = "[  ] Completed"
        set_cell_background(row_cells[0], "F8FAFC")
        set_cell_background(row_cells[1], "F8FAFC")
        
    doc.save(output_path)
    return output_path
